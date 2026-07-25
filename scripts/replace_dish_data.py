#!/usr/bin/env python3
"""Convert the cleaned recipe workbook into import-ready MySQL assets."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


EXPECTED_COLUMNS = [
    "来源ID",
    "菜名",
    "菜品大类",
    "菜系",
    "主要食材类型",
    "菜品形态",
    "烹饪方式",
    "口味",
    "用餐场景",
    "适用人群",
    "饮食属性",
    "设备",
    "难度",
    "预计时间分钟",
    "图片URL",
    "材料明细",
    "步骤数",
    "步骤",
]

OUTPUT_COLUMNS = [
    "id",
    "name",
    "type",
    "cl",
    "fl",
    "step",
    "tags",
    "image",
    "difficulty",
    "cook_time",
    "ingredients_amounts",
    "steps",
    "step_images",
    "tips",
    "methods",
    "kcal",
    "is_custom",
    "user_id",
    "create_time",
]

IMPORT_COLUMNS = [column for column in OUTPUT_COLUMNS if column != "id"]

CATEGORY_MAP = {
    "荤菜": "meat",
    "素菜": "veg",
    "汤羹": "soup",
    "甜品饮品": "dessert",
    "主食": "staple",
}

TAG_SOURCE_COLUMNS = [
    "菜系",
    "主要食材类型",
    "菜品形态",
    "口味",
    "用餐场景",
    "适用人群",
    "饮食属性",
    "设备",
]

TAG_SPLIT_RE = re.compile(r"[;,，、/|]+")


def text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def integer(value: Any, default: int = 0) -> int:
    if value is None or text(value) == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def source_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def unique_tags(row: dict[str, Any]) -> str:
    tags: list[str] = []
    seen: set[str] = set()
    for column in TAG_SOURCE_COLUMNS:
        for item in TAG_SPLIT_RE.split(text(row.get(column))):
            tag = item.strip()
            if tag and tag not in seen:
                seen.add(tag)
                tags.append(tag)
    return ",".join(tags)[:500]


def legacy_materials(raw: str) -> str:
    """Build the compact legacy CL field without duplicating the rich source."""
    ingredients: list[str] = []
    for item in raw.split("###"):
        parts = [part.strip() for part in item.split("|")]
        if not parts or not parts[0]:
            continue
        amount = parts[1] if len(parts) > 1 else ""
        unit = parts[2] if len(parts) > 2 else ""
        summary = f"{parts[0]}:{amount}{unit}".rstrip(":")
        ingredients.append(summary)
    return "#".join(ingredients)


def normalize_url(raw: str) -> str:
    if raw.startswith("http://"):
        return "https://" + raw[len("http://") :]
    return raw


def transform(row: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    issues: list[str] = []
    source_id = integer(row.get("来源ID"), -1)
    name = text(row.get("菜名"))
    source_category = text(row.get("菜品大类"))
    dish_type = CATEGORY_MAP.get(source_category, "")
    materials = text(row.get("材料明细"))
    steps = text(row.get("步骤"))
    difficulty = text(row.get("难度"))
    cook_minutes = integer(row.get("预计时间分钟"), 0)
    step_count = integer(row.get("步骤数"), 0)
    image = normalize_url(text(row.get("图片URL")))

    if source_id < 0:
        issues.append("invalid_source_id")
    if not name:
        issues.append("missing_name")
    if not dish_type:
        issues.append(f"unknown_category:{source_category}")
    if not materials:
        issues.append("missing_materials")
    if not steps:
        issues.append("missing_steps")
    if difficulty not in {"简单", "普通", "困难"}:
        issues.append(f"invalid_difficulty:{difficulty}")
    if cook_minutes <= 0:
        issues.append("invalid_cook_time")
    if step_count <= 0:
        issues.append("invalid_step_count")
    if image and not image.startswith("https://"):
        issues.append("invalid_image_url")

    output = {
        "id": source_id,
        "name": name,
        "type": dish_type,
        "cl": legacy_materials(materials),
        "fl": "2名成年人总量+15%冗余",
        "step": steps,
        "tags": unique_tags(row),
        "image": image,
        "difficulty": difficulty,
        "cook_time": f"{cook_minutes}分钟" if cook_minutes > 0 else "",
        "ingredients_amounts": materials,
        "steps": steps,
        "step_images": "[]",
        "tips": f"来源步骤数: {step_count}",
        "methods": text(row.get("烹饪方式")),
        "kcal": 0,
        "is_custom": 0,
        "user_id": "",
        "create_time": "",
    }
    return output, issues


def read_workbook(path: Path) -> tuple[str, list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if not workbook.sheetnames:
        raise ValueError("Workbook has no worksheets")
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(iterator)]
    missing = [column for column in EXPECTED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Workbook is missing columns: {', '.join(missing)}")

    rows: list[dict[str, Any]] = []
    for values in iterator:
        row = dict(zip(headers, values))
        if any(text(row.get(column)) for column in EXPECTED_COLUMNS):
            rows.append(row)
    return sheet.title, rows


def write_csv(path: Path, rows: Iterable[dict[str, Any]], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def sql_literal(value: Any) -> str:
    if value is None or value == "":
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    escaped = str(value).replace("\\", "\\\\").replace("'", "''")
    return f"'{escaped}'"


def write_sql(path: Path, rows: list[dict[str, Any]], source_name: str, batch_size: int) -> None:
    columns = OUTPUT_COLUMNS
    import_columns = IMPORT_COLUMNS
    with path.open("w", encoding="utf-8", newline="\n") as file:
        file.write("-- EatWhat system dish replacement generated from a cleaned workbook.\n")
        file.write(f"-- Source: {source_name}\n")
        file.write(f"-- Rows: {len(rows)}\n\n")
        file.write("SET NAMES utf8mb4;\n")
        file.write("SET time_zone = '+08:00';\n")
        file.write("USE food;\n\n")
        file.write("START TRANSACTION;\n")
        file.write("CREATE TEMPORARY TABLE eatwhat_food_import LIKE food;\n")
        file.write("ALTER TABLE eatwhat_food_import MODIFY COLUMN id INT NOT NULL;\n\n")

        for start in range(0, len(rows), batch_size):
            batch = rows[start : start + batch_size]
            file.write(f"INSERT INTO eatwhat_food_import ({', '.join(columns)}) VALUES\n")
            for index, row in enumerate(batch):
                values = ", ".join(sql_literal(row[column]) for column in columns)
                suffix = ",\n" if index < len(batch) - 1 else ";\n\n"
                file.write(f"({values}){suffix}")

        file.write("-- Keep user-created dishes and only replace system dishes.\n")
        file.write("DELETE FROM food WHERE user_id IS NULL;\n")
        file.write(f"INSERT INTO food ({', '.join(import_columns)})\n")
        file.write(f"SELECT {', '.join(import_columns)} FROM eatwhat_food_import;\n")
        file.write("DROP TEMPORARY TABLE eatwhat_food_import;\n")
        file.write("COMMIT;\n\n")
        file.write("-- Verification queries\n")
        file.write("SELECT type, COUNT(*) AS count FROM food WHERE user_id IS NULL GROUP BY type ORDER BY type;\n")
        file.write("SELECT COUNT(*) AS system_dishes FROM food WHERE user_id IS NULL;\n")
        file.write("SELECT COUNT(*) AS custom_dishes FROM food WHERE user_id IS NOT NULL;\n")


def write_report(
    path: Path,
    source: Path,
    sheet_name: str,
    source_rows: list[dict[str, Any]],
    output_rows: list[dict[str, Any]],
    issues: list[dict[str, Any]],
    source_display_name: str,
) -> None:
    categories = Counter(text(row.get("菜品大类")) for row in source_rows)
    output_types = Counter(row["type"] for row in output_rows)
    names: defaultdict[str, list[int]] = defaultdict(list)
    for row in output_rows:
        names[row["name"]].append(row["id"])
    duplicate_names = {name: ids for name, ids in names.items() if len(ids) > 1}
    missing = {
        column: sum(1 for row in source_rows if not text(row.get(column)))
        for column in EXPECTED_COLUMNS
    }

    lines = [
        "# 菜品数据替换审计报告",
        "",
        f"- 生成时间: `{datetime.now().astimezone().isoformat(timespec='seconds')}`",
        f"- 源文件: `{source_display_name}`",
        f"- 工作表: `{sheet_name}`",
        f"- SHA-256: `{source_hash(source)}`",
        f"- 源数据行: `{len(source_rows)}`",
        f"- 可导入行: `{len(output_rows)}`",
        f"- 阻断问题行: `{len(issues)}`",
        f"- 重复菜名组: `{len(duplicate_names)}`（来源ID不重复，因此保留全部版本）",
        "",
        "## 分类映射",
        "",
        "| 源分类 | 目标 type | 行数 |",
        "|---|---:|---:|",
    ]
    for source_category, target in CATEGORY_MAP.items():
        lines.append(f"| {source_category} | `{target}` | {categories[source_category]} |")

    lines.extend(["", "## 目标类型校验", "", "| type | 行数 |", "|---|---:|"])
    for dish_type in ("meat", "veg", "soup", "dessert", "staple"):
        lines.append(f"| `{dish_type}` | {output_types[dish_type]} |")

    lines.extend(["", "## 缺失情况", "", "仅列出存在缺失的源字段。", "", "| 字段 | 缺失行数 |", "|---|---:|"])
    for column, count in missing.items():
        if count:
            lines.append(f"| {column} | {count} |")

    lines.extend(
        [
            "",
            "## 替换策略",
            "",
            "- `来源ID` 写入暂存数据用于审计；导入正式 `food` 表时由 MySQL 重新分配自增 ID。",
            "- 这样可避免与现有用户自定义菜品 ID 冲突；已有日历记录仍可使用保存的 `dish_details` 展示历史内容。",
            "- 只删除 `user_id IS NULL` 的系统菜品，登录用户创建的自定义菜品保留。",
            "- 图片链接统一由 HTTP 转为 HTTPS。",
            "- `tags` 由菜系、食材类型、形态、口味、场景、人群、饮食属性和设备去重合并。",
            "- 原材料明细写入 `ingredients_amounts`，原步骤写入 `steps` 和兼容字段 `step`。",
            "- 源表没有热量数据，`kcal` 设置为 `0`，前端继续使用现有估算逻辑。",
            "- SQL 先写入临时表，再在事务中替换系统菜品；执行前仍必须做数据库备份。",
        ]
    )

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Cleaned recipe workbook path")
    parser.add_argument("--output-dir", required=True, help="Directory for generated assets")
    parser.add_argument("--batch-size", type=int, default=250, help="Rows per INSERT statement")
    parser.add_argument(
        "--source-display-name",
        help="Original source file name to preserve in the bundle and report",
    )
    args = parser.parse_args()

    source = Path(args.input).resolve()
    output_dir = Path(args.output_dir).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    output_dir.mkdir(parents=True, exist_ok=True)

    sheet_name, source_rows = read_workbook(source)
    output_rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    seen_ids: set[int] = set()
    for source_row in source_rows:
        output_row, row_issues = transform(source_row)
        if output_row["id"] in seen_ids:
            row_issues.append("duplicate_source_id")
        seen_ids.add(output_row["id"])
        if row_issues:
            issues.append(
                {
                    "来源ID": source_row.get("来源ID"),
                    "菜名": source_row.get("菜名"),
                    "问题": ";".join(row_issues),
                }
            )
        else:
            output_rows.append(output_row)

    source_display_name = args.source_display_name or source.name

    write_csv(output_dir / "food_import.csv", output_rows, OUTPUT_COLUMNS)
    write_csv(output_dir / "food_import_issues.csv", issues, ["来源ID", "菜名", "问题"])
    write_sql(output_dir / "replace_system_dishes.sql", output_rows, source_display_name, args.batch_size)
    write_report(
        output_dir / "dish_replacement_audit.md",
        source,
        sheet_name,
        source_rows,
        output_rows,
        issues,
        source_display_name,
    )
    shutil.copy2(source, output_dir / source_display_name)

    manifest = {
        "source": source_display_name,
        "source_sha256": source_hash(source),
        "sheet": sheet_name,
        "source_rows": len(source_rows),
        "import_rows": len(output_rows),
        "issue_rows": len(issues),
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "files": [
            source_display_name,
            "food_import.csv",
            "food_import_issues.csv",
            "replace_system_dishes.sql",
            "dish_replacement_audit.md",
        ],
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print(json.dumps(manifest, ensure_ascii=False))
    if issues:
        print("Blocking data issues found; review food_import_issues.csv", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
