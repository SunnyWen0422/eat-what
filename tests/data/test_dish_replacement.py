import csv
import hashlib
import importlib.util
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
BUNDLE = ROOT / "outputs" / "dish-replacement-20260718"
SCRIPT_PATH = ROOT / "scripts" / "replace_dish_data.py"
SPEC = importlib.util.spec_from_file_location("replace_dish_data", SCRIPT_PATH)
replacement = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(replacement)


def test_manifest_files_hash_and_counts_are_consistent():
    manifest = json.loads((BUNDLE / "manifest.json").read_text(encoding="utf-8"))
    for filename in manifest["files"]:
        assert (BUNDLE / filename).is_file(), filename

    source = BUNDLE / manifest["source"]
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    assert digest == manifest["source_sha256"]
    assert manifest["source_rows"] == 6665
    assert manifest["import_rows"] == 6665
    assert manifest["issue_rows"] == 0


def test_source_workbook_schema_dimensions_and_required_values():
    manifest = json.loads((BUNDLE / "manifest.json").read_text(encoding="utf-8"))
    workbook = load_workbook(BUNDLE / manifest["source"], read_only=True, data_only=True)
    assert workbook.sheetnames == [manifest["sheet"]]
    sheet = workbook[manifest["sheet"]]
    assert sheet.max_row == manifest["source_rows"] + 1
    assert sheet.max_column >= len(replacement.EXPECTED_COLUMNS)

    headers = [replacement.text(cell.value) for cell in next(sheet.iter_rows(max_row=1))]
    assert not set(replacement.EXPECTED_COLUMNS) - set(headers)


def test_import_csv_has_expected_type_distribution_and_no_blocking_gaps():
    with (BUNDLE / "food_import.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))

    assert len(rows) == 6665
    assert Counter(row["type"] for row in rows) == {
        "meat": 3200,
        "veg": 2190,
        "soup": 619,
        "dessert": 256,
        "staple": 400,
    }
    assert len({row["id"] for row in rows}) == len(rows)
    assert all(row["name"] and row["ingredients_amounts"] and row["steps"] for row in rows)
    assert all(not row["image"].startswith("http://") for row in rows)

    with (BUNDLE / "food_import_issues.csv").open(encoding="utf-8-sig", newline="") as handle:
        assert list(csv.DictReader(handle)) == []


def test_generated_sql_is_transactional_and_preserves_custom_dishes():
    sql = (BUNDLE / "replace_system_dishes.sql").read_text(encoding="utf-8")
    assert "START TRANSACTION;" in sql
    assert "CREATE TEMPORARY TABLE eatwhat_food_import LIKE food;" in sql
    assert "DELETE FROM food WHERE user_id IS NULL;" in sql
    assert "INSERT INTO food" in sql
    assert "COMMIT;" in sql
    assert sql.index("START TRANSACTION;") < sql.index("DELETE FROM food WHERE user_id IS NULL;") < sql.index("COMMIT;")


def test_generated_sql_preserves_utf8_recipe_text():
    sql = (BUNDLE / "replace_system_dishes.sql").read_text(encoding="utf-8")
    assert "青椒肉丝" in sql
    assert "丝瓜炒鸡蛋" in sql
    assert "闈掓鑲変笣" not in sql


def test_release_script_writes_mysql_backup_without_a_powershell_text_pipe():
    script = (ROOT / "scripts" / "apply_dish_replacement.ps1").read_text(encoding="utf-8")
    assert "--result-file=$backupFile" in script
    assert "$DbName | Set-Content" not in script


def test_transform_normalizes_url_tags_and_legacy_materials():
    row = {column: "" for column in replacement.EXPECTED_COLUMNS}
    row.update({
        "来源ID": 7,
        "菜名": "测试菜",
        "菜品大类": "荤菜",
        "难度": "简单",
        "预计时间分钟": 20,
        "图片URL": "http://example.com/a.jpg",
        "材料明细": "鸡蛋|2|个###盐|1|克",
        "步骤数": 2,
        "步骤": "第一步###第二步",
        "烹饪方式": "炒",
        "菜系": "家常",
        "口味": "家常",
    })

    result, issues = replacement.transform(row)
    assert issues == []
    assert result["type"] == "meat"
    assert result["image"] == "https://example.com/a.jpg"
    assert result["cl"] == "鸡蛋:2个#盐:1克"
    assert result["tags"] == "家常"
